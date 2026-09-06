"""Analysis workbook separated from immutable evidence, using the medical exporter style."""
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter

from .agent.exporter import AgentDatasetExportService


ANALYSIS_FIELDS = ["object_id", "mjd", "band", "magnitude", "magnitude_error",
    "magnitude_system", "time_unit", "time_scale", "redshift_heliocentric", "redshift_cmb",
    "quality_status", "quality_flags", "observation_id", "source_id", "source_row",
    "source_url", "metadata_match", "metadata_source_id", "metadata_source_url"]
LABELS = ["观测对象", "观测时间 MJD", "滤镜（保留原名）", "星等", "星等误差",
    "测光系统", "时间单位", "时间尺度", "日心红移", "CMB 红移", "观测质量状态", "质量提示",
    "观测编号", "来源编号", "原始行号", "来源网址", "元数据关联", "元数据来源编号", "元数据网址"]


def workbook_bytes(result, raw_tables):
    book = Workbook()
    book.remove(book.active)

    def table(name, headers, records):
        sheet = book.create_sheet(name[:31])
        sheet.append(headers)
        for record in records:
            sheet.append(record)
        # Catalog strings are evidence, never executable Excel formulas.
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
        AgentDatasetExportService._style_table(sheet, len(headers))
        sheet.freeze_panes = "B2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.sheet_view.showGridLines = False
        sheet.row_dimensions[1].height = 32
        for col in range(1, len(headers) + 1):
            sheet.column_dimensions[get_column_letter(col)].width = max(18, sheet.column_dimensions[get_column_letter(col)].width)
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                # Clip long raw text inside the cell, never spill into its neighbors.
                cell.alignment = Alignment(vertical="top", wrap_text=True)
            sheet.row_dimensions[row[0].row].height = 30
        return sheet

    cleaned = bool(result.get("pipeline_version"))
    report = result.get("cleaning_report", {})
    table("使用说明", ["项目", "说明"], [
        ["研究目标", result["topic"]], ["任务版本", result["task_id"]],
        ["上一个版本", result.get("parent_task_id", "首次运行")],
        ["处理状态", result.get("pipeline_version", "旧版仅字段解析；请在工作台执行清洗")],
        ["分析主表", "一行一次观测；数值是数值单元格，原始 JSON 不混入分析列"],
        ["原始证据", "原始来源表独立工作表；完整 raw_field/raw_value 在来源包 canonical_audit.json 中，按观测编号追溯"],
        ["质量状态", "usable_observation 仅表示通过观测字段检查；不表示校准、联合物理分析或发表已就绪"],
        ["复核方式", "待复核行不进入清洗主表。核对原始记录与来源后修正输入来源，再重新获取/清洗；本版不支持人工覆盖测量值"],
        *[["边界", item] for item in result["limitations"]]])
    sheet = table("清洗后观测" if cleaned else "标准化观测_未清洗", ANALYSIS_FIELDS,
                  [[r.get(k) for k in ANALYSIS_FIELDS] for r in result["rows"]])
    book.active = 1
    for row in sheet.iter_rows(min_row=2):
        for i in (1, 3, 4, 8, 9):
            row[i].number_format = "0.000000"
    table("字段字典", ["字段", "中文含义", "类型", "说明"], [
        [field, label, "数值" if field in {"mjd", "magnitude", "magnitude_error", "redshift_heliocentric", "redshift_cmb", "source_row"} else "文本",
         "空值表示来源未提供或未确认；不填零" if "redshift" in field else "source_native / unspecified 表示未确认统一校准口径" if field in {"magnitude_system", "time_scale"} else "按 observation_id / source_row 关联原始来源表"]
        for field, label in zip(ANALYSIS_FIELDS, LABELS)])
    names = {"input_rows": "输入观测", "usable_rows": "清洗后可用观测", "duplicate_rows": "完全重复", "invalid_rows": "无效隔离", "review_rows": "待复核",
             "numeric_values_converted": "文本转数值字段数", "whitespace_values_trimmed": "首尾空白规范数", "reconciled": "数量对账通过", "sorting": "排序规则",
             "imputed_values": "推算填补值数", "scientific_outlier_removals": "按科学离群规则删除数"}
    table("清洗报告", ["指标", "结果"], [[names.get(k, k), v] for k, v in report.items()] or [["状态", "旧结果未执行此清洗流程"]])
    table("质量检查", ["检查层", "判断", "检查依据", "证据"], [
        [g["label"], g["decision"], "；".join(g["checks"]), g["evidence"]] for g in result.get("quality_gate", {}).get("layers", [])])
    event_fields = ["observation_id", "rule", "field", "before", "after", "reason"]
    table("清洗记录", ["观测编号", "规则", "字段", "处理前", "处理后", "原因"],
          [[e.get(k) for k in event_fields] for e in result.get("cleaning_events", [])])
    table("隔离与复核", ["观测编号", "来源编号", "原始行号", "状态", "原因", "保留的重复观测编号"], [
        [r.get("observation_id", f"{r['source_id']}:{r['source_row']}"), r["source_id"], r["source_row"], status, r["reason"], r.get("duplicate_of")]
        for key, status in (("rejected_rows", "无效"), ("review_rows", "待复核"), ("duplicate_rows", "完全重复")) for r in result.get(key, [])])
    table("数据来源", ["来源编号", "来源网址", "获取时间", "SHA256", "原始文件"], [
        [s.get(k) for k in ("source_id", "url", "retrieved_at", "sha256", "file")] for s in result["sources"]])
    for name, source_id, fields, records in raw_tables:
        table("原始_" + name, ["source_id", "source_row", *fields], [[source_id, i, *[r.get(k) for k in fields]] for i, r in enumerate(records, 1)])
    stream = io.BytesIO()
    book.save(stream)
    return stream.getvalue()
